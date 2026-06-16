#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generates a detailed PPTX report of the manual-vs-optimal comparison (TTS),
in ENGLISH, including each module's full name and how the decision variables
are distributed over the Pareto.

Usage:
    python generar_reporte_ppt.py --json comparacion_manual/comparacion_manual.json \
        --out report_comparison.pptx
"""
import argparse, json, os
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# ── Full module names ─────────────────────────────────────────────────────────
NAMES = {
    "M4":  "SMAC Black-Box — Bayesian Optimization (Gaussian Process + EI)",
    "M4RF":"SMAC Random Forest (HPO: Random Forest + EI + Sobol)",
    "M7":  "SMAC + Stochastic Kriging (EI)",
    "M8":  "Adaptive Stochastic Kriging",
    "M10": "Stochastic Kriging — KGCP (Knowledge Gradient for Continuous Parameters)",
    "M11": "ASTRO-DF (Adaptive Sampling Trust-Region Optimization, derivative-free)",
    "M13": "SPSA (Simultaneous Perturbation Stochastic Approximation)",
    "RS":  "Random Search",
}
# Decision variables: (min, max, type) for normalization + English label
VARIABLES = {
    "horas_especialista_1ra":   (8, 30, "int",   "First-consult slots/week"),
    "horas_control_post":       (20, 70, "int",  "Post-control slots/week"),
    "cupos_laboratorio_ugd":    (20, 100, "int", "Lab slots (UGD)"),
    "cupos_ecografia_matrona":  (10, 50, "int",  "Ultrasound slots (midwife)"),
    "cupos_ecografia_ugd":      (10, 50, "int",  "Ultrasound slots (UGD)"),
    "dias_publicacion":         (1, 10, "int",   "Scheduling lead time (days)"),
    "num_matronas":             (1, 4, "int",    "# Midwives"),
    "num_agentes_ugd":          (1, 4, "int",    "# UGD agents"),
    "pct_bloqueo_1ra":          (0.05, 0.5, "float", "First-consult blocking %"),
    "pct_consultas_vacias":     (0.05, 0.5, "float", "Empty-consult %"),
    "pct_no_contactabilidad":   (0.05, 0.5, "float", "Non-contactability %"),
    "pct_bloqueo_post_control": (0.05, 0.5, "float", "Post-control blocking %"),
}
BLUE, RED = "#1f77b4", "#d62728"


def disp(nombre):
    """Display name: 'ÓPTIMO M8' → 'Model M8'."""
    return nombre.replace("ÓPTIMO", "Model").strip()


def color_of(tipo): return BLUE if tipo == "optimo" else RED
def label_var(v):   return VARIABLES[v][3]


# ──────────────────────────────────────────────────────────────────────────────
# FIGURES
# ──────────────────────────────────────────────────────────────────────────────
def fig_tts_bars(esc, out):
    e = sorted(esc, key=lambda x: x["tts_media"])
    nom = [disp(x["escenario"]) for x in e]
    m = [x["tts_media"] for x in e]
    lo = [x["tts_media"] - x["tts_ic95"][0] for x in e]
    hi = [x["tts_ic95"][1] - x["tts_media"] for x in e]
    col = [color_of(x["tipo"]) for x in e]
    plt.figure(figsize=(11, 5.5))
    plt.bar(nom, m, yerr=[lo, hi], capsize=4, color=col, alpha=0.85,
            edgecolor="k", linewidth=0.5)
    for i, v in enumerate(m):
        plt.text(i, v + max(hi) + 2, f"{v:.0f}", ha="center", fontsize=9)
    plt.ylabel("TTS — mean time in system [days]  (lower = better)")
    plt.title("TTS by option (mean ± 95% CI)")
    plt.xticks(rotation=25, ha="right")
    plt.grid(axis="y", alpha=0.3)
    import matplotlib.patches as mp
    plt.legend(handles=[mp.Patch(color=BLUE, label="Optimal (model)"),
                        mp.Patch(color=RED, label="Manual")], loc="upper left")
    plt.tight_layout(); p = os.path.join(out, "fig_tts_bars.png")
    plt.savefig(p, dpi=150); plt.close(); return p


def fig_pareto(esc, out):
    plt.figure(figsize=(10, 6.5))
    for x in esc:
        at = x["atenciones_media"]
        if not np.isfinite(at): continue
        c = color_of(x["tipo"]); mk = "o" if x["tipo"] == "optimo" else "s"
        plt.scatter(at, x["tts_media"], s=120, marker=mk, color=c, zorder=3,
                    edgecolor="k", linewidth=0.7)
        plt.annotate(disp(x["escenario"]), (at, x["tts_media"]), fontsize=9,
                     xytext=(6, 4), textcoords="offset points")
    import matplotlib.patches as mp
    plt.legend(handles=[mp.Patch(color=BLUE, label="Optimal (model)"),
                        mp.Patch(color=RED, label="Manual")])
    plt.xlabel("Patients served (total attentions)  →  more = better")
    plt.ylabel("TTS — time in system [days]  ←  less = better")
    plt.title("Observed results (TTS vs throughput)\n"
              "Note: only TTS was optimized; throughput is an emergent outcome",
              fontsize=11)
    plt.grid(alpha=0.3)
    plt.annotate("better", xy=(0.97, 0.04), xycoords="axes fraction",
                 fontsize=11, color="green", ha="right",
                 bbox=dict(boxstyle="round", fc="honeydew", ec="green"))
    plt.tight_layout(); p = os.path.join(out, "fig_pareto.png")
    plt.savefig(p, dpi=150); plt.close(); return p


def fig_pareto_variables(esc, vd, out):
    """Pareto colored by 4 key decision variables → how they distribute."""
    keys = ["dias_publicacion", "pct_bloqueo_1ra", "num_agentes_ugd", "horas_control_post"]
    fig, axes = plt.subplots(2, 2, figsize=(12, 9))
    pts = [(x["escenario"], x["atenciones_media"], x["tts_media"])
           for x in esc if np.isfinite(x["atenciones_media"])]
    for ax, var in zip(axes.ravel(), keys):
        lo, hi, _, lab = VARIABLES[var]
        xs = [p[1] for p in pts]; ys = [p[2] for p in pts]
        vals = [vd.get(p[0], {}).get(var, np.nan) for p in pts]
        sc = ax.scatter(xs, ys, c=vals, s=170, cmap="viridis", vmin=lo, vmax=hi,
                        edgecolor="k", linewidth=0.6, zorder=3)
        for (nom, a, t), v in zip(pts, vals):
            ax.annotate(f"{v:.2f}" if isinstance(v, float) and v < 5 else f"{v:.0f}",
                        (a, t), fontsize=7, ha="center", va="center", color="white")
        cb = fig.colorbar(sc, ax=ax, fraction=0.046); cb.set_label(lab, fontsize=8)
        ax.set_xlabel("patients served"); ax.set_ylabel("TTS [days]")
        ax.set_title(f"Pareto colored by: {lab}", fontsize=10)
        ax.grid(alpha=0.3)
    fig.suptitle("How decision variables distribute over the Pareto", fontsize=13)
    fig.tight_layout(rect=[0, 0, 1, 0.97])
    p = os.path.join(out, "fig_pareto_variables.png")
    fig.savefig(p, dpi=150); plt.close(); return p


def fig_parallel(esc, vd, out):
    """Parallel coordinates of the 12 (normalized) variables per option."""
    options = [x["escenario"] for x in esc]
    tipos = {x["escenario"]: x["tipo"] for x in esc}
    varnames = list(VARIABLES)
    fig, ax = plt.subplots(figsize=(13, 6.5))
    xpos = np.arange(len(varnames))
    for nom in options:
        dd = vd.get(nom, {})
        ys = []
        for v in varnames:
            lo, hi, _, _ = VARIABLES[v]
            val = dd.get(v)
            ys.append((val - lo) / (hi - lo) if val is not None and hi > lo else np.nan)
        ax.plot(xpos, ys, marker="o", ms=4, lw=1.8, alpha=0.85, color=color_of(tipos[nom]))
        ax.annotate(disp(nom), (xpos[-1], ys[-1]), fontsize=7, xytext=(4, 0),
                    textcoords="offset points", va="center")
    ax.set_xticks(xpos)
    ax.set_xticklabels([label_var(v) for v in varnames], rotation=35, ha="right", fontsize=8)
    ax.set_ylabel("normalized position in admissible range (0 = min, 1 = max)")
    ax.set_title("Decision variables by option (parallel coordinates)\n"
                 "blue = optimal (models)   ·   red = manual", fontsize=11)
    ax.grid(alpha=0.3)
    fig.tight_layout(); p = os.path.join(out, "fig_parallel.png")
    fig.savefig(p, dpi=150); plt.close(); return p


# ──────────────────────────────────────────────────────────────────────────────
# PPTX helpers
# ──────────────────────────────────────────────────────────────────────────────
def _title(slide, text, sub=None):
    tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.95))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text
    p.font.size = Pt(26); p.font.bold = True; p.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    if sub:
        p2 = tf.add_paragraph(); p2.text = sub
        p2.font.size = Pt(13); p2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def _bullets(slide, items, left=0.5, top=1.3, w=12.3, h=5.6, size=14):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, (txt, lvl) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt; p.level = lvl
        p.font.size = Pt(size - 2*lvl)
        if lvl == 0: p.font.bold = True


def _table(slide, data, left, top, w, h, header=True, fontsize=10):
    rows, cols = len(data), len(data[0])
    t = slide.shapes.add_table(rows, cols, Inches(left), Inches(top),
                               Inches(w), Inches(h)).table
    for r in range(rows):
        for c in range(cols):
            cell = t.cell(r, c); cell.text = str(data[r][c])
            for para in cell.text_frame.paragraphs:
                para.font.size = Pt(fontsize)
                if header and r == 0:
                    para.font.bold = True
                    para.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if header and r == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    return t


def _img(slide, path, left, top, w=None, h=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    slide.shapes.add_picture(path, Inches(left), Inches(top), **kw)


def blank(prs): return prs.slides.add_slide(prs.slide_layouts[6])


# ──────────────────────────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--json", default="comparacion_manual/comparacion_manual.json")
    ap.add_argument("--out", default="report_comparison.pptx")
    ap.add_argument("--figdir", default="reporte_figs")
    args = ap.parse_args()
    os.makedirs(args.figdir, exist_ok=True)

    d = json.load(open(args.json, encoding="utf-8"))
    esc = d["escenarios"]; av = d["analisis_varianza"]
    rec = d["recomendacion"]; vd = d["variables_decision"]
    n_man = sum(1 for x in esc if x["tipo"] == "manual")
    n_opt = sum(1 for x in esc if x["tipo"] == "optimo")
    n_rep_man = len(esc[0]["tts_raw"])

    f_bar = fig_tts_bars(esc, args.figdir)
    f_par = fig_pareto(esc, args.figdir)
    f_pv  = fig_pareto_variables(esc, vd, args.figdir)
    f_pc  = fig_parallel(esc, vd, args.figdir)

    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)

    # 1) Title
    s = blank(prs)
    _title(s, "Clinic scheduling optimization (CRS)",
           "Manual adjustments vs. optimization models  ·  "
           "Indicator: TTS = mean patient time in system [days]")
    _bullets(s, [
        ("Objective: minimize total patient time in system (TTS).", 0),
        (f"Configurations compared: {n_man} manual + {n_opt} optimal (models).", 1),
        (f"Replications: manual n={n_rep_man} (re-simulated); models n=50 (stored re-evaluation).", 1),
        ("Tests: Welch t-test (unequal variances) + Welch ANOVA + Kruskal-Wallis + Levene.", 1),
        ("Run-size (MODULE 1): recommended 10-18 replications; baseline TTS 95% CI = [254.7, 266.3] d.", 1),
    ], top=2.6, size=16)

    # 2) Methods / modules with names
    s = blank(prs)
    _title(s, "Optimization methods (modules)")
    rows = [["Key", "Full method name"]]
    for k in ["M4", "M4RF", "M7", "M8", "M10", "M11", "M13", "RS"]:
        rows.append([k, NAMES[k]])
    _table(s, rows, 0.5, 1.4, 12.3, 4.6, fontsize=12)
    _bullets(s, [("This batch ran: M4, M7, M8, M10, M13 (15 seeds each). "
                  "M4RF and M11 are pending the PC run. RS produced no valid data.", 0)],
             top=6.2, size=12)

    # 3) TTS results (table + bars)
    s = blank(prs)
    _title(s, "Results — TTS per option (mean ± SD, 95% CI, Welch vs Current)")
    best_tts = min(x["tts_media"] for x in esc)
    rows = [["Option", "type", "TTS μ±σ", "95% CI", "served", "Δ vs Cur", "p Welch", "verdict"]]
    for x in sorted(esc, key=lambda x: x["tts_media"]):
        at = x["atenciones_media"]; p = x["p_welch_vs_current"]
        verdict = "—" if x["escenario"] == "Current" else (
            "BETTER" if (np.isfinite(p) and p < 0.05) else "≈")
        rows.append([disp(x["escenario"]), "optimal" if x["tipo"] == "optimo" else "manual",
                     f"{x['tts_media']:.1f}±{x['tts_sd']:.1f}",
                     f"[{x['tts_ic95'][0]:.1f},{x['tts_ic95'][1]:.1f}]",
                     f"{at:.0f}" if np.isfinite(at) else "—",
                     f"{x['delta_tts_vs_current']:+.1f}",
                     f"{p:.1e}" if np.isfinite(p) else "—", verdict])
    _table(s, rows, 0.3, 1.35, 6.4, 5.0, fontsize=9)
    _img(s, f_bar, 6.95, 1.5, w=6.15)

    # 4) Analysis of variance
    s = blank(prs)
    _title(s, "Analysis of variance (robust to heteroscedasticity)")
    wa = av["welch_anova"]
    _bullets(s, [
        ("Levene test (homogeneity of variances):", 0),
        (f"p = {av['levene_p']:.1e}  →  variances are UNEQUAL (heteroscedastic) → use Welch.", 1),
        ("Welch ANOVA (omnibus, unequal variances):", 0),
        (f"F({wa['df1']:.0f}, {wa['df2']:.1f}) = {wa['F']:.1f},  p = {wa['p']:.1e}  → strong differences.", 1),
        ("Kruskal-Wallis (non-parametric backup):", 0),
        (f"p = {av['kruskal_wallis_p']:.1e}.", 1),
        ("Post-hoc Welch t + Holm correction:", 0),
        ("ALL optimal models are significantly better than ALL manual scenarios (p ≪ 0.001).", 1),
        ("Current differs significantly from Management, Mgmt+Cap and v2.", 1),
    ], top=1.4, size=15)

    # 5) Business decision
    s = blank(prs)
    _title(s, "Business decision — for the decision-maker")
    cur = rec["current"]; base_t, base_a = cur["tts"], cur["atenciones"]
    rows = [["Option", "wait [d]", "Δwait", "served", "Δserved", "signif.", "Pareto"]]
    for x in sorted(esc, key=lambda x: x["tts_media"]):
        at = x["atenciones_media"]
        dt = x["tts_media"] - base_t
        da = at - base_a if np.isfinite(at) else np.nan
        p = x["p_welch_vs_current"]
        sig = "—" if x["escenario"] == "Current" else ("yes" if (np.isfinite(p) and p < 0.05) else "no")
        par = "★" if x["escenario"] in rec["pareto"] else ""
        rows.append([disp(x["escenario"]), f"{x['tts_media']:.0f}",
                     f"{dt:+.0f} ({100*dt/base_t:+.0f}%)",
                     f"{at:.0f}" if np.isfinite(at) else "—",
                     f"{da:+.0f}" if np.isfinite(da) else "—", sig, par])
    _table(s, rows, 0.3, 1.35, 9.0, 4.7, fontsize=10)
    _bullets(s, [
        (f"▸ Lowest wait: {disp(rec['menor_tts'])}.", 0),
        (f"▸ Highest throughput: {disp(rec['mayor_throughput'])}.", 0),
        ("▸ Models cut waiting by ~33% vs ~9% for the best manual.", 0),
    ], left=9.5, top=1.5, w=3.6, size=12)

    # 6) Pareto + caveat
    s = blank(prs)
    _title(s, "Observed results: TTS vs throughput")
    _img(s, f_par, 0.4, 1.3, w=8.2)
    _bullets(s, [
        ("Critical reading:", 0),
        ("Only TTS was optimized by the models.", 1),
        ("Throughput (attentions) is an EMERGENT outcome, not an objective.", 1),
        ("Therefore no hypothesis test on throughput; it is context only.", 1),
        ("Message: reducing waiting did NOT sacrifice throughput (it rose).", 1),
        ("To optimize both, use f = TTS − λ·attentions.", 1),
    ], left=8.8, top=1.4, w=4.3, size=13)

    # 7) Decision variables (parallel coordinates)
    s = blank(prs)
    _title(s, "Decision variables by option — what each configuration changes")
    _img(s, f_pc, 0.3, 1.35, w=12.7)
    _bullets(s, [("Models use levers the manual scenarios did not: lead time = 1 day (not 7), "
                  "blocking at 5% (not 10%), 4 agents, post-control hours = 70 — without adding midwives.", 0)],
             top=6.6, size=12)

    # 8) Decision variables OVER the Pareto (requested)
    s = blank(prs)
    _title(s, "How decision variables distribute over the Pareto")
    _img(s, f_pv, 0.3, 1.3, w=8.6)
    _bullets(s, [
        ("Each panel colors the points by a key decision variable.", 0),
        ("The low-TTS region (bottom) concentrates:", 0),
        ("Scheduling lead time = 1 day (immediate publishing).", 1),
        ("First-consult blocking = 5% (minimal blocking).", 1),
        ("# UGD agents = 4 (maximum agent staffing).", 1),
        ("Post-control slots = 70 (maximum).", 1),
        ("→ These are the levers separating models from manual scenarios.", 0),
    ], left=9.1, top=1.4, w=4.0, size=12)

    # 9) Conclusions
    s = blank(prs)
    _title(s, "Conclusions and recommendation")
    _bullets(s, [
        ("1. Optimization models dominate manual adjustments on TTS.", 0),
        (f"Best model: {disp(rec['menor_tts'])} (~33% less waiting), significant (p ≪ 0.001).", 1),
        ("2. The improvement does not sacrifice throughput (attentions increase).", 0),
        ("3. Key levers found by the models: immediate publishing, minimal blocking, "
         "agent staffing — rather than adding midwives.", 0),
        ("4. Validity: same simulator for all; Welch test (heteroscedastic); 95% CI; "
         "replication count supported by MODULE 1.", 0),
        ("5. Next step: CRN r=50 run (gold standard); to co-optimize time + throughput, "
         "use the combined objective f = TTS − λ·attentions.", 0),
    ], top=1.4, size=15)

    prs.save(args.out)
    n = len(prs.slides._sldIdLst)
    print(f"OK report generated: {args.out}  ({n} slides)")


if __name__ == "__main__":
    main()
